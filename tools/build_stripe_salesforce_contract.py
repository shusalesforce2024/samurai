from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path

OUT = Path('docs/contracts/Stripe_CSV_Salesforce投入_業務委託契約書_ドラフト.docx')
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY='17365D'; BLUE='2E74B5'; PALE='E8EEF5'; GRAY='666666'; LIGHT='F2F4F7'; BLACK='000000'

def font(run, size=10.5, bold=False, color=BLACK):
    run.font.name='Yu Gothic'; run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Yu Gothic')
    run._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); run._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri')
    run.font.size=Pt(size); run.bold=bold; run.font.color.rgb=RGBColor.from_string(color)

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)

def margins(cell, top=80, start=120, bottom=80, end=120):
    tc=cell._tc.get_or_add_tcPr(); tcMar=tc.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tc.append(tcMar)
    for tag,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        el=tcMar.find(qn('w:'+tag))
        if el is None: el=OxmlElement('w:'+tag); tcMar.append(el)
        el.set(qn('w:w'),str(val)); el.set(qn('w:type'),'dxa')

def set_repeat(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)

def table(rows, widths, header=True, sizes=None):
    t=doc.add_table(rows=0, cols=len(widths)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    tblPr=t._tbl.tblPr; tblW=tblPr.find(qn('w:tblW')); tblW.set(qn('w:w'),'9360'); tblW.set(qn('w:type'),'dxa')
    ind=OxmlElement('w:tblInd'); ind.set(qn('w:w'),'120'); ind.set(qn('w:type'),'dxa'); tblPr.append(ind)
    grid=t._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for w in widths:
        gc=OxmlElement('w:gridCol'); gc.set(qn('w:w'),str(w)); grid.append(gc)
    for ri,data in enumerate(rows):
        cells=t.add_row().cells
        if ri==0 and header: set_repeat(t.rows[-1])
        for ci,(cell,text) in enumerate(zip(cells,data)):
            cell.width=Inches(widths[ci]/1440); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(cell)
            tcW=cell._tc.get_or_add_tcPr().find(qn('w:tcW')); tcW.set(qn('w:w'),str(widths[ci])); tcW.set(qn('w:type'),'dxa')
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
            r=p.add_run(str(text)); font(r, sizes[ci] if sizes else 8.5, ri==0 and header)
            if ri==0 and header: shade(cell,PALE)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def para(text='', bold=False, align=None, after=5, indent=0, size=10.5):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.1
    if indent: p.paragraph_format.first_line_indent=Inches(indent)
    if align is not None: p.alignment=align
    font(p.add_run(text),size,bold); return p

def heading(text, level=1):
    p=doc.add_paragraph(style=f'Heading {level}'); p.add_run(text); return p

def bullet(text):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(3); font(p.add_run(text),10); return p

def new_page(): doc.add_page_break()

doc=Document()
sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1); sec.header_distance=sec.footer_distance=Inches(.492)
styles=doc.styles
normal=styles['Normal']; normal.font.name='Yu Gothic'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Yu Gothic'); normal.font.size=Pt(10.5); normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.1
for name,size,before,after,color in [('Heading 1',16,14,8,BLUE),('Heading 2',13,11,6,BLUE),('Heading 3',11.5,8,4,NAVY)]:
    s=styles[name]; s.font.name='Yu Gothic'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'Yu Gothic'); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True
for name in ['List Bullet','List Number']:
    s=styles[name]; s.font.name='Yu Gothic'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'Yu Gothic'); s.font.size=Pt(10); s.paragraph_format.left_indent=Inches(.5); s.paragraph_format.first_line_indent=Inches(-.25); s.paragraph_format.space_after=Pt(3)

# header/footer
h=sec.header.paragraphs[0]; h.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(h.add_run('Stripe CSV / Salesforce投入　業務委託契約書（ドラフト）'),8.5,False,GRAY)
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
font(f.add_run('本書は当事者確認用ドラフトです　｜　'),8,False,GRAY)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)

# cover
para('業務委託契約書',True,WD_ALIGN_PARAGRAPH.CENTER,8,size=24)
para('Stripe CSVのSalesforce投入・MRR把握基盤整備',False,WD_ALIGN_PARAGRAPH.CENTER,28,size=14)
table([
    ['項目','内容'],['委託者（甲）','［会社名・住所・代表者名を記入］'],['受託者（乙）','［会社名・住所・代表者名を記入］'],['契約形態','準委任型の業務委託'],['委託料','154,000円（税込）'],['想定工数','28時間'],['契約期間','［開始日］から［終了日］まで'],['作成日','2026年8月20日'],['版','ドラフト 1.0']], [2700,6660], sizes=[9,9.5])
para('重要事項',True,after=4,size=11)
para('本書の［　］部分は締結前に当事者間で確定する。法務・税務上の最終確認は、各当事者の専門家による確認を推奨する。',False,after=0,size=9.5)
new_page()

heading('業務委託契約条項',1)
para('［委託者名］（以下「甲」という。）と［受託者名］（以下「乙」という。）は、Stripeから出力したCSVデータをSalesforceへ投入し、商品別・月別MRRを把握するための基盤整備業務について、次のとおり業務委託契約（以下「本契約」という。）を締結する。')

heading('第1条（目的）',2)
para('本契約は、Stripe由来の顧客、サブスクリプション、請求及び決済情報を、外部IDを用いてSalesforceへ安全かつ重複なく取り込み、Stripeと請求情報を合わせた商品別・月別MRRをSalesforce上で把握できる状態を整備することを目的とする。')

heading('第2条（契約の性質）',2)
para('本契約は準委任契約とし、乙は善良な管理者の注意をもって本業務を遂行する。乙は、特定の経営成果、売上、MRR数値、又は第三者サービスの継続稼働を保証するものではない。')

heading('第3条（委託業務）',2)
para('甲は乙に、別紙1「業務仕様書」に定める設計、Salesforce設定、Excel変換・検証テンプレート作成、Data Loaderマッピング及び手順書作成、Sandbox試験並びに照合作業を委託し、乙はこれを受託する。')
para('業務上の基本方式は月1回の手動運用とし、Stripe DashboardからのCSV出力、Excelテンプレートによる変換・検証、Data Loaderによる外部ID Upsert、Salesforce上での件数・金額・エラー確認の順で実施する。')

heading('第4条（成果物）',2)
for x in ['Salesforce追加項目・選択リスト・権限・レイアウトの設定又は設定内容一覧','Stripe CSV用Excel変換・検証テンプレート','Data Loader Upsert用マッピングファイル及び月次運用手順書','Sandbox試験結果、異常系試験結果及び件数・金額照合結果','例外一覧の出力仕様（商品未確定、金額不一致、複数候補その他）']:
    bullet(x)
para('成果物のファイル形式、保管場所及び引渡方法は、甲乙協議のうえ定める。Salesforce本番環境への投入は、甲が書面又は電磁的方法により明示的に承認した場合に限り実施する。')

heading('第5条（契約月次明細の取扱い）',2)
para('商品別・月別MRRを請求データのみから正確かつ継続的に算出できることがSandbox試験及び照合により確認できた場合、甲乙は、ContractPeriod__c又はContractLineItem__c等の契約月次情報の新規投入を省略できる。省略可否は、対象商品の特性、解約・数量変更・日割り・割引・追加課金・返金調整の有無及び既存MRR管理ロジックへの影響を確認したうえで、甲が承認する。')

heading('第6条（作業期間及び進行）',2)
para('本業務の作業期間は［開始日］から［終了日］までとする。甲による資料提供、環境利用許可、確認又は承認の遅延、第三者サービスの障害その他乙の責めに帰さない事由がある場合、乙は甲と協議のうえ期限を合理的に変更できる。')

heading('第7条（甲の協力事項）',2)
for x in ['Stripe及びSalesforceの必要な権限、テスト環境、対象CSV、商品・Price対応情報の提供','取引先の名寄せ方針、商品マスタの正解データ、MRR算定基準及び検収担当者の指定','個人情報を含むデータの取扱方針、バックアップ及び本番投入承認','成果物又は質問事項に対する合理的な期間内の回答']:
    bullet(x)

heading('第8条（委託料及び支払）',2)
para('本業務の委託料は154,000円（税込）とする。乙は［請求時期］に請求書を発行し、甲は請求書受領月の［　］日までに乙指定口座へ振り込む。振込手数料は甲の負担とする。')
para('前項の金額は想定28時間の別紙業務を対象とする。甲の要請による仕様変更、対象データの著しい増加、データ品質問題への追加対応、第三者サービス仕様変更その他当初想定を超える作業は、第9条に従う。')

heading('第9条（追加作業及び変更管理）',2)
para('Invoice Lines API取得ツール、Webhook、Salesforce Apexによる常時API連携、専用ログ・移行Work・決済履歴オブジェクト、Contact又はOpportunityの自動作成その他別紙1で初期対象外とした事項は、別途見積り及び合意を要する。Invoice Lines API取得ツールの追加工数目安は8時間から16時間とするが、金額及び納期は着手前の書面又は電磁的方法による合意で確定する。')
para('変更依頼は、変更内容、理由、工数、金額、納期及び影響範囲を確認し、甲乙の承認後に実施する。')

heading('第10条（検収）',2)
para('乙は成果物の完成後、甲に引き渡し、甲は引渡日から5営業日以内に別紙2「検収基準」に基づき確認する。不適合がある場合、甲は当該期間内に具体的内容を通知し、乙は本契約範囲内で合理的な修正を行う。期間内に通知がない場合、成果物は検収済みとみなす。ただし、隠れた不適合については第16条を適用する。')

heading('第11条（データ管理及び安全対策）',2)
for x in ['Insertではなく、原則としてStripe IDを外部IDとしたUpsertを使用する。','本番投入前に対象データをバックアップし、入力、成功及びエラーCSVを月別に保存する。','取引先の曖昧一致は自動反映せず、商品未確定、金額不一致又は複数候補のデータは除外し、例外一覧に出力する。','Customer別及びInvoice別に件数と金額を照合し、UTC日時を日本時間へ変換する。','本番環境に対する操作は最小権限で行い、甲の承認した手順及び対象に限定する。']:
    bullet(x)

heading('第12条（管理元及び既存処理との関係）',2)
para('Salesforceを契約、請求及びMRRの管理元とし、Stripeをサブスクリプション状態及び決済結果の管理元とする。請求管理元がStripeである契約は、既存のSalesforce更新請求及びFreee自動作成の対象から除外する。既存の契約月次明細バッチは、MRR管理との整合性を確認したうえで継続する。')

heading('第13条（秘密保持・個人情報）',2)
para('甲及び乙は、本業務に関連して知り得た相手方の技術上、営業上その他一切の非公知情報を秘密として取り扱い、本業務以外に使用せず、相手方の事前承諾なく第三者へ開示しない。ただし、法令に基づく開示、公知情報、受領前から保有していた情報又は正当な権限を有する第三者から取得した情報を除く。')
para('乙は、個人情報を取り扱う場合、個人情報保護法その他関係法令を遵守し、目的達成に必要な範囲に限定して取り扱う。本条の義務は本契約終了後3年間存続し、個人情報に関する義務は法令上必要な期間存続する。')

heading('第14条（知的財産権）',2)
para('本業務のために新たに作成された成果物の著作権は、委託料の完済を条件として甲に移転する。ただし、乙又は第三者が従前から保有するプログラム、テンプレート、ノウハウ、汎用部品及びオープンソースソフトウェア等の権利は移転しない。乙は、甲が成果物を利用するために必要な範囲で、これらを非独占的かつ無期限に利用することを許諾する。')

heading('第15条（第三者サービス）',2)
para('Stripe、Salesforce、Data Loader、Microsoft Excelその他第三者サービスの仕様変更、停止、障害、利用制限又はデータ不備に起因する事象について、乙は自己の責めに帰すべき場合を除き責任を負わない。各サービスの利用契約及び利用料は甲の責任と負担による。')

heading('第16条（契約不適合への対応）',2)
para('検収後30日以内に、本契約又は別紙仕様に適合しない事項が判明し、かつ乙の責めに帰すべき場合、乙は無償で合理的な修正を行う。第三者仕様変更、甲の操作、提供データの誤り又は合意済み仕様の変更は無償修正の対象外とする。')

heading('第17条（損害賠償）',2)
para('甲又は乙が本契約に違反し相手方に損害を与えた場合、通常かつ直接の損害に限り賠償責任を負う。乙の賠償額の総額は、乙の故意又は重過失、秘密保持違反若しくは個人情報漏えいの場合を除き、本契約に基づき甲が乙に支払った委託料総額を上限とする。逸失利益、間接損害及び特別損害は賠償対象外とする。')

heading('第18条（解除）',2)
para('一方当事者が本契約に重大な違反をし、相当期間を定めた催告後も是正しない場合、相手方は本契約を解除できる。支払停止、破産等の申立て、事業廃止又は信用状態の重大な悪化があった場合は、催告なく解除できる。解除までに完了した作業及び合理的に発生した費用について、甲は乙に支払う。')

heading('第19条（反社会的勢力の排除）',2)
para('甲及び乙は、自ら及び役員等が反社会的勢力に該当せず、これらと関係を有しないことを表明し保証する。違反が判明した場合、相手方は催告なく本契約を解除できる。')

heading('第20条（不可抗力）',2)
para('天災、感染症、戦争、法令変更、通信障害、クラウドサービスの広域障害その他合理的支配を超える事由により義務の履行が遅延又は不能となった場合、当該当事者はその責任を負わず、速やかに相手方へ通知し対応を協議する。')

heading('第21条（再委託）',2)
para('乙は、本業務の全部又は重要な一部を第三者に再委託する場合、甲の事前承諾を得る。乙は再委託先に本契約と同等の義務を負わせ、その履行について責任を負う。')

heading('第22条（権利義務の譲渡禁止）',2)
para('甲及び乙は、相手方の事前の書面による承諾なく、本契約上の地位又は権利義務を第三者へ譲渡し、又は担保に供してはならない。')

heading('第23条（協議・準拠法・管轄）',2)
para('本契約に定めのない事項又は解釈上の疑義は、甲乙誠実に協議して解決する。本契約は日本法に準拠し、本契約に関する紛争の第一審専属的合意管轄裁判所は［東京地方裁判所／当事者間で合意する裁判所］とする。')

heading('署名欄',1)
para('本契約成立の証として、本書2通を作成し、甲乙記名押印のうえ各1通を保有する。電磁的記録により締結する場合、各当事者は当該電磁的記録を保管する。')
table([['締結日','［　年　月　日］'],['甲','住所：［　　　　　　　　　　　　　　　］\n会社名：［　　　　　　　　　　　　　　］\n代表者：［　　　　　　　　　　　　　］　印'],['乙','住所：［　　　　　　　　　　　　　　　］\n会社名：［　　　　　　　　　　　　　　］\n代表者：［　　　　　　　　　　　　　］　印']], [1800,7560], header=False, sizes=[9.5,10])

new_page(); heading('別紙1　業務仕様書',1)
heading('1. 基本方針',2)
for x in ['月1回の運用とし、Webhook及び常時API連携は初期対象外とする。','Stripe CSVをExcelテンプレートで変換・検証し、Data LoaderでUpsertする。','Stripe IDをSalesforceの外部IDとして保持し、重複作成を防止する。','Salesforceを契約・請求・MRRの管理元とし、Stripeをサブスクリプション状態・決済結果の管理元とする。']:
    bullet(x)
heading('2. 対象オブジェクトとマッピング方針',2)
table([
['Salesforce対象','Stripe元','取得','判定・条件'],
['取引先 Account','Customer','CSV可','Customer ID外部ID追加後にUpsert'],
['取引先責任者 Contact','Customer・申込者','一部','初期は自動作成しない。担当者氏名がある場合も別途判断'],
['商品マスタ ProductMaster__c','Product・Price','API推奨','Price ID対応表により一意に特定'],
['取引 Opportunity__c','metadata・WEB申込','不足','初期は自動作成しない。受注チャネル等が必要'],
['契約管理 Contract__c','Subscription','CSV可','Subscription ID外部ID追加後にUpsert'],
['契約期間 ContractPeriod__c','Current Period','CSV可','現在期間は作成可。MRR目的上の要否を試験で判定'],
['契約月次明細 ContractLineItem__c','Subscription Item','不足','商品確定後に作成可。請求のみでMRR把握可能なら省略可'],
['請求 Invoice__c','Invoice','CSV可','Invoice ID外部ID追加後にUpsert'],
['請求明細 InvoiceLine__c','Invoice Line Item','API取得','条件内はSubscription CSVから生成。複雑ケースはAPI追加'],
['Stripe決済','Transaction・Charge','CSV可','初期は請求既存項目へ最新状態のみ反映']], [1900,1600,1100,4760], sizes=[7.6,7.6,7.6,7.6])

heading('3. 最小追加項目',2)
table([
['オブジェクト','追加内容','用途'],
['Account','StripeCustomerId__c','Customer外部ID・重複防止'],
['Contract__c','StripeSubscriptionId__c','Subscription外部ID'],
['ProductMaster__c','StripePriceId__c','Stripe Priceと商品の対応'],
['Invoice__c','StripeInvoiceId__c','Invoice外部ID'],
['Invoice__c','StripeChargeId__c','最新決済の追跡'],
['Contract__c','請求管理元に「Stripe」追加','Freee更新請求から除外'],
['Contract__c','作成元に「StripeImport」追加','CSV取込契約の識別']], [2100,2900,4360], sizes=[8.2,8.2,8.2])

heading('4. 月次取得データ',2)
table([
['データ','頻度','用途'],['Customers CSV','毎月','取引先の新規・更新・名寄せ'],['Subscriptions CSV','毎月','契約の開始・継続・解約・数量変更'],['Invoices CSV','毎月','請求金額、請求日、状態'],['Transactions CSV','毎月','決済状態、入金額、入金日、返金額'],['Product・Price対応表','初回・商品変更時','Price IDと商品マスタの対応']], [2600,1700,5060], sizes=[8.5,8.5,8.5])

heading('5. Data Loader投入順序',2)
for x in ['商品マスタをStripePriceId__cでUpsert','取引先をStripeCustomerId__cでUpsert','契約管理をStripeSubscriptionId__cでUpsert','必要と判定された場合、新規契約の契約期間及び初回契約月次明細を投入','請求をStripeInvoiceId__cでUpsert','Transactions CSVから決済状態、入金額及び入金日を更新','成功・エラーCSVとSalesforceの件数・金額を照合']:
    bullet(x)

heading('6. 請求明細の生成条件',2)
para('次の条件をすべて満たす場合に限り、Subscription CSVから請求明細を生成する。')
for x in ['Subscription Itemが1種類','日割りがない','割引・クーポンがない','追加課金・返金調整がない','単価×数量とInvoice金額が一致','Stripe Price IDから商品マスタが一意に決まる']:
    bullet(x)
para('条件を満たさない請求は例外一覧へ出力する。例外件数・金額・発生パターンを甲乙で確認し、運用上許容できない場合に限り、Invoice Lines API取得ツールを追加作業として検討する。')

heading('7. 初期対象外',2)
for x in ['Stripe Webhook','Salesforce ApexによるStripe API常時連携','Stripeイベントオブジェクト','Stripe連携ログオブジェクト','Stripe移行Workオブジェクト群','決済履歴専用オブジェクト','Contactの自動作成','Opportunityの自動作成']:
    bullet(x)
para('決済履歴は初期段階では請求の既存項目に最新状態のみ保持する。')

heading('8. 工数内訳',2)
table([
['作業','工数'],['外部ID・選択リスト・権限・レイアウト','6時間'],['Excel変換・検証テンプレート','8時間'],['Data Loaderマッピング・手順書','6時間'],['Sandbox試験・異常系・照合','8時間'],['合計','28時間']], [7000,2360], sizes=[9,9])
para('Invoice Lines API取得ツールを追加する場合の工数目安：追加8～16時間（別途見積り・合意）。',bold=True,size=9.5)

new_page(); heading('別紙2　検収基準',1)
table([
['No.','検収項目','合格基準'],
['1','外部ID・重複防止','指定した外部IDで再投入しても同一レコードが更新され、意図しない重複が作成されない。'],
['2','商品対応','Stripe Price IDから商品マスタが一意に決まり、未確定・複数候補は除外される。'],
['3','契約・請求Upsert','対象CSVから契約及び請求が仕様どおり作成又は更新される。'],
['4','決済状態反映','最新の決済状態、入金額、入金日及び必要な返金情報が請求へ反映される。'],
['5','MRR把握','合意した対象月について、商品別・月別MRRを算出・確認できる。契約月次明細を省略する場合は請求ベースの算出結果が照合される。'],
['6','例外制御','曖昧一致、商品未確定、金額不一致、複数候補及び複雑な請求が自動反映されず、例外一覧に出力される。'],
['7','日時変換','対象UTC日時が合意したルールで日本時間へ変換される。'],
['8','照合','Customer別・Invoice別の件数及び金額が、入力CSV、Data Loader結果及びSalesforceで照合できる。'],
['9','証跡','入力、成功、エラー及び例外CSVが月別に保存され、手順書から追跡できる。'],
['10','既存処理影響','請求管理元=Stripeの契約が既存Salesforce更新請求及びFreee自動作成から除外されることをSandboxで確認できる。']], [600,2900,5860], sizes=[7.7,7.7,7.7])
heading('検収時の前提・未確定事項',2)
for x in ['本番環境の項目API名、選択リスト値、既存バッチ及び権限構成は着手時に再確認する。','商品別・月別MRRの具体的な算式、対象期間、税込・税抜、返金・値引き・日割りの扱いは甲が承認する。','契約期間・契約月次明細の投入要否は、請求データのみでMRR目的を満たせるかを試験して決定する。','月次の処理対象件数及び許容例外率は［　］件、［　］％又は甲乙が別途合意する基準とする。','本番投入作業を成果物に含めるか、甲が手順書に基づき実施するかを着手前に確定する。']:
    bullet(x)

heading('別紙3　月次運用フロー（概要）',1)
for i,x in enumerate(['Stripe DashboardからCustomers、Subscriptions、Invoices、Transactions CSVを出力','Excel変換・検証テンプレートへ取込み、形式・必須値・外部ID・商品・金額・日時を検証','商品、取引先、契約、必要な契約期間／月次明細、請求の順に外部IDUpsert','Transactions CSVから最新決済情報を請求へ更新','入力・成功・エラー・例外CSVとSalesforceをCustomer別・Invoice別に照合','処理結果、差異、例外対応及び承認者を記録し、月別フォルダへ保存'],1):
    p=doc.add_paragraph(style='List Number'); p.paragraph_format.space_after=Pt(5); font(p.add_run(x),10)

# core props
doc.core_properties.title='Stripe CSV Salesforce投入 業務委託契約書（ドラフト）'
doc.core_properties.subject='商品別・月別MRR把握基盤整備'
doc.core_properties.author=''
doc.core_properties.keywords='Stripe, Salesforce, CSV, Data Loader, MRR, 業務委託契約'
doc.save(OUT)
print(OUT.resolve())
