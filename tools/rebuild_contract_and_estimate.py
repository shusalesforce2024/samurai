from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('docs/contracts')
OUT.mkdir(parents=True, exist_ok=True)
CONTRACT = OUT / 'Stripe_CSV_Salesforce投入_業務委託契約書_ドラフト.docx'
ESTIMATE = OUT / 'Stripe_CSV_Salesforce投入_見積書_ドラフト.docx'

def set_font(run, size=10.5, bold=False):
    run.font.name = 'Yu Gothic'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Yu Gothic')
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    run.font.size = Pt(size)
    run.bold = bold

def setup(title):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.top_margin, s.bottom_margin = Mm(18), Mm(18)
    s.left_margin, s.right_margin = Mm(22), Mm(22)
    normal = d.styles['Normal']
    normal.font.name = 'Yu Gothic'; normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Yu Gothic')
    normal.font.size = Pt(10.5); normal.paragraph_format.space_after = Pt(3); normal.paragraph_format.line_spacing = 1.05
    for st in ('List Bullet', 'List Number'):
        d.styles[st].font.name = 'Yu Gothic'; d.styles[st]._element.rPr.rFonts.set(qn('w:eastAsia'), 'Yu Gothic')
        d.styles[st].font.size = Pt(10); d.styles[st].paragraph_format.space_after = Pt(2)
    d.core_properties.title = title; d.core_properties.author = ''
    return d

def p(d, text='', size=10.5, bold=False, align=None, after=3, first=0):
    x = d.add_paragraph(); x.paragraph_format.space_after = Pt(after); x.paragraph_format.line_spacing = 1.05
    if first: x.paragraph_format.first_line_indent = Mm(first)
    if align is not None: x.alignment = align
    set_font(x.add_run(text), size, bold)
    return x

def clause(d, title, bodies=None, bullets=None):
    p(d, title, 11, True, after=2)
    for body in bodies or []: p(d, body, first=4)
    for item in bullets or []:
        x=d.add_paragraph(style='List Bullet'); x.paragraph_format.left_indent=Mm(8); x.paragraph_format.first_line_indent=Mm(-3); set_font(x.add_run(item),10)

def set_cell(cell, text, width_twips, bold=False, align=None, fill=None, size=9):
    cell.width = Mm(width_twips / 56.6929)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn('w:tcW')); tcW.set(qn('w:w'), str(width_twips)); tcW.set(qn('w:type'), 'dxa')
    mar=OxmlElement('w:tcMar')
    for side,val in [('top',90),('start',100),('bottom',90),('end',100)]:
        e=OxmlElement('w:'+side); e.set(qn('w:w'),str(val)); e.set(qn('w:type'),'dxa'); mar.append(e)
    tcPr.append(mar)
    if fill:
        shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
    par=cell.paragraphs[0]; par.paragraph_format.space_after=Pt(0)
    if align is not None: par.alignment=align
    set_font(par.add_run(str(text)),size,bold)

def tbl(d, rows, widths, header=True, sizes=None):
    t=d.add_table(rows=0, cols=len(widths)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    grid=t._tbl.tblGrid
    for c in list(grid): grid.remove(c)
    for w in widths:
        g=OxmlElement('w:gridCol'); g.set(qn('w:w'),str(w)); grid.append(g)
    pr=t._tbl.tblPr; tw=pr.find(qn('w:tblW')); tw.set(qn('w:w'),str(sum(widths))); tw.set(qn('w:type'),'dxa')
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        if ri==0 and header:
            h=OxmlElement('w:tblHeader'); h.set(qn('w:val'),'true'); t.rows[-1]._tr.get_or_add_trPr().append(h)
        for ci,val in enumerate(row):
            set_cell(cells[ci],val,widths[ci],ri==0 and header,None,'E7E6E6' if ri==0 and header else None,(sizes or [9]*len(widths))[ci])
    p(d,'',after=1)
    return t

# ---------------- Contract: follows the supplied reference's concise 18-article structure ----------------
d=setup('業務委託契約書')
p(d,'業務委託契約書',18,True,WD_ALIGN_PARAGRAPH.CENTER,14)
p(d,'本契約は、以下の当事者間において締結される。')
for item in ['甲：株式会社SAMURAI ARCHITECTS（以下「甲」という）','乙：川波嵩（以下「乙」という）']:
    x=d.add_paragraph(style='List Bullet'); set_font(x.add_run(item),10.5)

clause(d,'第1条（目的）',[ '本契約は、甲の営業DX推進を目的として、乙がStripeから出力したCSVデータのSalesforce投入および商品別・月別MRR把握基盤の整備に関する業務を受託し、その実施条件を定めることを目的とする。'])
clause(d,'第2条（業務内容）',[ '乙は、別途提示する見積書に基づき、以下の業務（以下「本業務」という）を実施する。','1. 業務範囲'],[
    'Stripe CSVとSalesforceオブジェクトのマッピングおよび要件整理',
    'Salesforce外部ID、選択リスト、権限およびレイアウトの設定',
    'Excel変換・検証テンプレートの作成',
    'Data Loader Upsert用マッピングおよび月次運用手順書の作成',
    'Sandboxにおける正常系・異常系試験および件数・金額照合',
    'Stripe由来の請求情報と商品情報による商品別・月別MRR把握方法の整備'])
p(d,'2. 成果物',bold=True)
for item in ['Salesforce設定または設定内容一覧','Excel変換・検証テンプレート','Data Loaderマッピングおよび月次運用手順書','Sandbox試験・照合結果','例外一覧の出力仕様']:
    x=d.add_paragraph(style='List Bullet'); set_font(x.add_run(item),10)
p(d,'※業務範囲、成果物、前提条件および対象外事項の詳細は、別途提示する見積書による。',9.5,after=4)

clause(d,'第3条（契約形態）',['本契約は準委任契約とし、乙は善良なる管理者の注意義務をもって本業務を遂行する。なお、本契約は特定の成果の完成、売上、MRR数値または第三者サービスの継続稼働を保証するものではない。'])
clause(d,'第4条（開発方式）',['本業務は、月1回のCSV取込運用を前提とし、WBSまたは双方が合意した作業計画に基づき実施する。業務内容および優先順位は、甲乙協議の上、柔軟に変更できるものとする。','契約月次明細の投入要否は、請求情報のみで商品別・月別MRRを正確かつ継続的に把握できるかをSandboxで確認し、甲乙協議の上で決定する。'])
clause(d,'第5条（契約期間）',['開始日：2026年8月24日','終了日：2026年9月30日'])
clause(d,'第6条（報酬および支払条件）',['1. 報酬　本業務の報酬は以下とする。'],['総額：154,000円（税込）','内訳および詳細：別途提示する見積書の内容に基づく'])
p(d,'2. 支払条件　甲は乙に対し、以下の条件で支払う。')
for item in ['契約締結時（着手金）：77,000円（税込）','業務完了後：77,000円（税込）','支払期日：各請求月末締め翌月末払い']:
    x=d.add_paragraph(style='List Bullet'); set_font(x.add_run(item),10)
p(d,'業務完了は、WBSまたは合意した作業計画に基づく対象作業の完了をいう。',9.5,first=4)

clause(d,'第7条（再委託）',['乙は、本業務の全部または一部を第三者に再委託する場合、事前に甲の書面による承認を得るものとする。乙は再委託先の行為について責任を負い、再委託先に本契約と同等の秘密保持義務および情報セキュリティ義務を課す。'])
clause(d,'第8条（資料提供）',['甲は、乙が本業務を遂行するために必要なStripeおよびSalesforceの情報、CSV、商品・Price対応情報、アカウント、権限その他の資料を適時提供するものとする。'])
clause(d,'第9条（変更管理）',['業務内容は進行に応じて変更されることを前提とする。ただし、当初想定を超える変更、データ品質問題への追加対応または初期対象外事項への対応が必要となる場合、双方協議の上、追加費用および期間を決定する。Invoice Lines API取得ツールを追加する場合の想定追加工数は8時間から16時間とし、着手前に別途見積りおよび合意を行う。'])
clause(d,'第10条（知的財産権）',['1. 本業務により新たに作成された成果物の知的財産権は、報酬の完済を条件として甲に帰属する。','2. 乙または第三者が従来から保有するノウハウ、テンプレート、汎用部品その他の権利は乙または当該第三者に帰属する。','3. 乙は、甲を特定できないよう匿名化した形で本業務実績を利用できる。'])
clause(d,'第11条（秘密保持）',['双方は、本契約に関連して知り得た相手方の秘密情報および個人情報を、本業務以外に使用せず、相手方の事前承諾なく第三者に開示してはならない。本条は契約終了後も有効とする。'])
clause(d,'第12条（責任範囲）',['1. 乙の損害賠償責任は、乙の故意または重過失による場合を除き、本契約に基づき受領した報酬額を上限とする。','2. 逸失利益その他の間接損害または特別損害について、乙は責任を負わない。'])
clause(d,'第13条（免責事項）',['以下に起因する損害または不利益について、乙は自己の責めに帰すべき場合を除き責任を負わない。'],['Salesforce、Stripe、Data LoaderまたはExcelの仕様変更','API制限または外部サービス障害','甲が提供したデータの不備','甲による設定変更、操作または業務運用','第三者の行為'])
clause(d,'第14条（契約解除）',['1. 当事者の一方が本契約に違反し、相当期間を定めた催告後も是正されない場合、相手方は本契約を解除できる。','2. 甲は30日前の通知により本契約を中途解約できる。','3. 中途解約の場合、乙は実施済み業務分および合理的に発生した費用を請求できる。'])
clause(d,'第15条（不具合対応）',['乙は、本業務の完了後30日間に限り、本業務に起因する不具合について合理的な範囲で修正対応を行う。ただし、甲による設定・運用変更、第三者サービスの仕様変更、提供データの不備または本業務範囲外の機能に関する事項は対象外とする。'])
clause(d,'第16条（反社会的勢力の排除）',['双方は、自らおよび役員等が反社会的勢力に該当せず、これらと関係を有しないことを保証し、違反した場合、相手方は本契約を直ちに解除できる。'])
clause(d,'第17条（協議事項）',['本契約に定めのない事項または解釈上の疑義については、双方誠意をもって協議し解決する。'])
clause(d,'第18条（管轄）',['本契約に関する紛争は、東京地方裁判所を第一審の専属的合意管轄裁判所とする。'])

p(d,'署名',12,True,after=3)
p(d,'本契約締結の証として、本書2通を作成し、双方記名押印の上、各1通を保有する。電磁的記録により締結する場合、各当事者は当該電磁的記録を保管する。')
p(d,'2026年8月24日',align=WD_ALIGN_PARAGRAPH.RIGHT,after=8)
p(d,'甲',bold=True); p(d,'東京都港区南青山3丁目4番7号 第7SYビル201',first=4); p(d,'株式会社SAMURAI ARCHITECTS',first=4); p(d,'代表取締役　加藤利基　印',first=4,after=9)
p(d,'乙',bold=True); p(d,'東京都江東区大島1丁目23-22',first=4); p(d,'ラポール西大島ウエスト408',first=4); p(d,'川波　嵩　印',first=4)
d.save(CONTRACT)

# ---------------- Estimate: separated former appendices ----------------
e=setup('見積書')
p(e,'御　見　積　書',20,True,WD_ALIGN_PARAGRAPH.CENTER,10)
p(e,'株式会社SAMURAI ARCHITECTS　御中',11,True,after=6)
tbl(e,[['見積日','2026年8月20日'],['見積番号','［　　　　　　　　　］'],['有効期限','見積日より30日間'],['件名','Stripe CSVのSalesforce投入・MRR把握基盤整備']], [1700,7460], header=False, sizes=[9,9.5])
p(e,'下記のとおりお見積り申し上げます。',after=5)
tbl(e,[['御見積金額（税込）','154,000円']], [4500,4660], header=False, sizes=[11,14])

p(e,'1. 見積明細',12,True,after=4)
tbl(e,[['No.','作業項目','内容','工数','金額（税込）'],
['1','Salesforce設定','外部ID、選択リスト、権限、レイアウト','6h','33,000円'],
['2','Excelテンプレート','Stripe CSVの変換・入力検証・例外抽出','8h','44,000円'],
['3','Data Loader・手順書','Upsertマッピング、投入順序、月次運用手順','6h','33,000円'],
['4','Sandbox試験・照合','正常系・異常系、件数・金額、既存処理影響確認','8h','44,000円'],
['','合計','','28h','154,000円']], [600,1900,4200,900,1560], sizes=[7.5,8,8,8,8])
p(e,'※単価換算：5,500円／時間（税込）。上記金額には消費税を含みます。',9,after=6)

p(e,'2. 作業目的',12,True)
p(e,'Stripe由来の顧客・契約・請求・決済情報をSalesforceへ安全に投入し、Stripeと請求情報を合わせた商品別・月別MRRを把握できる状態を整備します。請求データのみで目的を達成できる場合、契約期間・契約月次明細の新規投入は省略します。')
p(e,'3. 基本方式',12,True)
for item in ['月1回、Stripe DashboardからCSVを出力','Excel変換・検証テンプレートで形式、商品、金額、日時および例外を確認','Stripe IDを外部IDとしてData LoaderでUpsert','Salesforceで件数、金額およびエラーを照合','Salesforceを契約・請求・MRRの管理元、Stripeをサブスクリプション状態・決済結果の管理元とする']:
    x=e.add_paragraph(style='List Bullet'); set_font(x.add_run(item),10)

p(e,'4. 対象オブジェクト・追加項目',12,True)
tbl(e,[['Salesforce対象','Stripe元','対応方針'],
['Account','Customer','StripeCustomerId__cを外部IDとしてUpsert'],
['ProductMaster__c','Product・Price','StripePriceId__cで商品を一意に対応'],
['Contract__c','Subscription','StripeSubscriptionId__cでUpsert。請求管理元=Stripe、作成元=StripeImport'],
['ContractPeriod__c','Current Period','必要性を試験で判断。請求のみでMRR把握可能なら省略'],
['ContractLineItem__c','Subscription Item','商品確定後に作成可。MRR目的上不要なら省略'],
['Invoice__c','Invoice・Transaction','StripeInvoiceId__cでUpsert。StripeChargeId__cと最新決済状態を保持'],
['InvoiceLine__c','Invoice Line Item','単純請求のみSubscription CSVから生成。複雑請求は例外化']], [2300,1900,4960], sizes=[8,8,8])

p(e,'5. 月次投入順序',12,True)
for item in ['商品マスタをStripePriceId__cでUpsert','取引先をStripeCustomerId__cでUpsert','契約管理をStripeSubscriptionId__cでUpsert','必要な場合のみ契約期間・初回契約月次明細を投入','請求をStripeInvoiceId__cでUpsert','Transactions CSVから決済状態、入金額および入金日を更新','成功・エラーCSVとSalesforceの件数・金額を照合']:
    x=e.add_paragraph(style='List Number'); set_font(x.add_run(item),10)

p(e,'6. 請求明細の生成条件',12,True)
p(e,'次の条件をすべて満たす場合のみ、Subscription CSVから請求明細を生成します。')
for item in ['Subscription Itemが1種類','日割り、割引・クーポン、追加課金、返金調整がない','単価×数量とInvoice金額が一致','Stripe Price IDから商品マスタが一意に決まる']:
    x=e.add_paragraph(style='List Bullet'); set_font(x.add_run(item),10)
p(e,'条件を満たさない請求は例外一覧へ出力します。例外が運用上許容できない場合、Invoice Lines API取得ツールを追加見積りします。',9.5)

p(e,'7. 安全対策・照合',12,True)
for item in ['Insertではなく外部IDUpsertを使用','本番投入前に対象データをバックアップ','取引先の曖昧一致は自動反映しない','商品未確定、金額不一致、複数候補は除外して例外一覧へ出力','Customer別・Invoice別に件数と金額を照合','UTC日時を日本時間へ変換','入力・成功・エラーCSVを月別に保存','請求管理元=Stripeの契約を既存Salesforce更新請求・Freee自動作成から除外']:
    x=e.add_paragraph(style='List Bullet'); set_font(x.add_run(item),10)

p(e,'8. 初期対象外・追加費用',12,True)
for item in ['Stripe Webhook','Salesforce ApexによるStripe API常時連携','Stripeイベント・連携ログ・移行Workオブジェクト','決済履歴専用オブジェクト','Contact・Opportunityの自動作成','本番投入作業（別途明示合意がある場合を除く）']:
    x=e.add_paragraph(style='List Bullet'); set_font(x.add_run(item),10)
tbl(e,[['追加候補','想定追加工数','条件'],['Invoice Lines API取得ツール','8～16時間','例外件数・金額・発生パターン確認後に別途見積り・合意']], [3600,1700,3860], sizes=[8.5,8.5,8.5])

p(e,'9. 前提条件・支払条件',12,True)
for item in ['甲から必要なStripe CSV、商品・Price対応情報、Salesforce権限および確認担当者が提供されること','契約期間および作業日程は契約締結前に確定すること','支払条件：契約締結時77,000円、業務完了後77,000円（各請求月末締め翌月末払い）','仕様変更、対象データの著しい増加、データ品質問題への追加対応は変更管理の対象','成果物：Salesforce設定または設定一覧、Excelテンプレート、Data Loaderマッピング・手順書、Sandbox試験・照合結果、例外一覧仕様']:
    x=e.add_paragraph(style='List Bullet'); set_font(x.add_run(item),10)

p(e,'発行者',11,True,after=3)
p(e,'川波　嵩',bold=True); p(e,'東京都江東区大島1丁目23-22'); p(e,'ラポール西大島ウエスト408')
e.save(ESTIMATE)

print(CONTRACT.resolve())
print(ESTIMATE.resolve())
